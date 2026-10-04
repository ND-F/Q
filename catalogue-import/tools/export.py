import json, csv, segno
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
D=json.load(open('rows.json')); COLS=D['cols']; EXTRA=D['extra']; rows=D['rows']
OUT='/home/user/Q/catalogue-import/'
link_col=get_column_letter(COLS.index('LINK')+1)
def qr_formula(i):
    return f'=IMAGE("https://api.qrserver.com/v1/create-qr-code/?size=300x300&margin=8&color=06444C&bgcolor=EAE8D8&data="&ENCODEURL({link_col}{i}))'
wb=Workbook(); ws=wb.active; ws.title='new pieces'
hdr=COLS+EXTRA; ws.append(hdr)
for i,r in enumerate(rows, start=2):
    vals=[r[c] for c in hdr]; vals[COLS.index('QR')]=qr_formula(i); ws.append(vals)
for c in range(1,len(hdr)+1):
    cell=ws.cell(1,c); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='06444C' if c<=len(COLS) else '8A7A4F')
    ws.column_dimensions[get_column_letter(c)].width=18
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment=Alignment(wrap_text=True, vertical='top')
ws.freeze_panes='C2'
ws2=wb.create_sheet('skipped'); ws2.append(['inventory','reason'])
for s in D['skipped']: ws2.append(list(s))
wb.save(OUT+'mia-catalogue-rows.xlsx')
with open(OUT+'mia-catalogue-rows.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(hdr)
    for i,r in enumerate(rows, start=2):
        vals=[r[c] for c in hdr]; vals[COLS.index('QR')]=qr_formula(i); w.writerow(vals)
# QR بهوية نديم: نقط تيل على بيج، وعيون الزوايا عنابى (زى qr.html)
import cv2
det=cv2.QRCodeDetector()
kw=dict(scale=12, border=3, dark='#06444C', light='#EAE8D8', finder_dark='#4A121C', finder_light='#EAE8D8')
for r in rows:
    # اتأكد إن الكود بيتقري؛ لو لأ جرّب mask تانى
    for mask in [None,0,1,2,3,4,6,7]:
        q=segno.make(r['LINK'], error='h', mask=mask)
        png=OUT+f"qr/q{r['id']}.png"; q.save(png, **kw)
        if det.detectAndDecode(cv2.imread(png))[0]==r['LINK']: break
    else: raise SystemExit('QR not readable: '+r['id'])
    q.save(OUT+f"qr/q{r['id']}.svg", **kw)
print(len(rows),'rows;', 'LINK column', link_col)
